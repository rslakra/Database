"""Excel file connector — requires optional openpyxl dependency."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from dataio.base import ConnectorError, DataConnector
from datatype.data_formats import RowList


class ExcelConnector(DataConnector):
    format_name = "excel"
    file_extensions = ("xlsx", "xlsm")

    def read(self, path: str | Path, **options: Any) -> RowList:
        openpyxl = _require_openpyxl()
        path = Path(path)
        if not path.exists():
            raise ConnectorError(f"File not found: {path}")

        sheet_name = options.get("sheet_name")
        data_only = options.get("data_only", True)

        wb = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
        try:
            ws = wb[sheet_name] if sheet_name else wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                headers = next(rows_iter)
            except StopIteration:
                return []
            if not headers:
                return []

            fieldnames = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(headers)]
            result: RowList = []
            for row in rows_iter:
                if row is None or all(cell is None for cell in row):
                    continue
                result.append(
                    {
                        fieldnames[i]: row[i] if i < len(row) else None
                        for i in range(len(fieldnames))
                    }
                )
            return result
        finally:
            wb.close()

    def write(self, path: str | Path, rows: RowList, **options: Any) -> None:
        openpyxl = _require_openpyxl()
        if not rows:
            raise ConnectorError("Cannot write Excel with zero rows")

        path = Path(path)
        sheet_name = options.get("sheet_name", "Sheet1")
        fieldnames = options.get("fieldnames") or list(rows[0].keys())

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(col) for col in fieldnames])

        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        wb.close()


def _require_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise ConnectorError(
            "Excel support requires openpyxl. Install with: pip install openpyxl"
        ) from exc
    return openpyxl
